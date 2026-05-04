"""
Excel logger — three sheets, date-wise analytics, embedded screenshots.

Sheets:
  "All Detections"      — every vehicle exit, every day
  "Should Install"      — vehicles with sticker_count < 3  (need installation)
  "Unknown Stickers"    — vehicles with unrecognised stickers (screenshot embedded)
"""

import threading
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Optional

import cv2
import numpy as np
import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from PIL import Image as PILImage
    _PIL = True
except ImportError:
    _PIL = False


# ── Sheet / column definitions ────────────────────────────────────────────────
SH_ALL     = "All Detections"
SH_INSTALL = "Should Install"
SH_UNKNOWN = "Unknown Stickers"

_ALL_HDR     = ["Date", "Time", "Vehicle_Number", "Logo_Triangle", "AM_MOTORS",
                "Website_Sticker", "Stickers_Found"]
_INSTALL_HDR = ["Date", "Time", "Vehicle_Number", "Logo_Triangle", "AM_MOTORS",
                "Website_Sticker", "Stickers_Missing", "Required_Action"]
_UNKNOWN_HDR = ["Date", "Time", "Vehicle_Number", "Notes", "Screenshot"]

_HDR_FONT = Font(bold=True, color="FFFFFF")
_FILLS = {
    "all":     PatternFill("solid", fgColor="1F497D"),
    "install": PatternFill("solid", fgColor="7B2C00"),
    "unknown": PatternFill("solid", fgColor="3D1A78"),
    3: PatternFill("solid", fgColor="C6EFCE"),
    2: PatternFill("solid", fgColor="FFEB9C"),
    1: PatternFill("solid", fgColor="FFE0B2"),
    0: PatternFill("solid", fgColor="FFC7CE"),
}
_WIDTHS = {
    SH_ALL:     [14, 10, 16, 14, 12, 15, 13],
    SH_INSTALL: [14, 10, 16, 14, 12, 15, 14, 35],
    SH_UNKNOWN: [14, 10, 16, 35, 22],
}


def _setup_sheet(ws, headers: list, fill_key: str, widths: list):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = _HDR_FONT
        cell.fill = _FILLS[fill_key]
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _frame_to_xl_image(frame: np.ndarray, size=(180, 130)) -> Optional[XLImage]:
    if not _PIL or frame is None:
        return None
    try:
        rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        pil = PILImage.fromarray(rgb)
        pil.thumbnail(size, PILImage.LANCZOS)
        buf = BytesIO()
        pil.save(buf, format="PNG")
        buf.seek(0)
        return XLImage(buf)
    except Exception:
        return None


# ── Main logger class ─────────────────────────────────────────────────────────
class ExcelLogger:
    def __init__(self, filepath: str):
        self.filepath = Path(filepath)
        self._lock = threading.Lock()
        self._init_file()

    def _init_file(self):
        self.filepath.parent.mkdir(parents=True, exist_ok=True)
        if self.filepath.exists():
            # Ensure all three sheets exist in old files
            try:
                wb = openpyxl.load_workbook(str(self.filepath))
                dirty = False
                if SH_ALL not in wb.sheetnames:
                    ws = wb.create_sheet(SH_ALL, 0)
                    _setup_sheet(ws, _ALL_HDR, "all", _WIDTHS[SH_ALL])
                    dirty = True
                if SH_INSTALL not in wb.sheetnames:
                    ws = wb.create_sheet(SH_INSTALL)
                    _setup_sheet(ws, _INSTALL_HDR, "install", _WIDTHS[SH_INSTALL])
                    dirty = True
                if SH_UNKNOWN not in wb.sheetnames:
                    ws = wb.create_sheet(SH_UNKNOWN)
                    _setup_sheet(ws, _UNKNOWN_HDR, "unknown", _WIDTHS[SH_UNKNOWN])
                    dirty = True
                if dirty:
                    wb.save(str(self.filepath))
            except Exception:
                pass
            return

        wb = Workbook()
        # First sheet
        wb.active.title = SH_ALL
        _setup_sheet(wb[SH_ALL], _ALL_HDR, "all", _WIDTHS[SH_ALL])
        ws2 = wb.create_sheet(SH_INSTALL)
        _setup_sheet(ws2, _INSTALL_HDR, "install", _WIDTHS[SH_INSTALL])
        ws3 = wb.create_sheet(SH_UNKNOWN)
        _setup_sheet(ws3, _UNKNOWN_HDR, "unknown", _WIDTHS[SH_UNKNOWN])
        wb.save(str(self.filepath))

    # ── Logging ───────────────────────────────────────────────────────────────

    def log(
        self,
        vehicle_number: str,
        logo1: bool,
        logo2: bool,
        logo3: bool,
    ) -> bool:
        count = sum([logo1, logo2, logo3])
        now = datetime.now()
        date_s = now.strftime("%Y-%m-%d")
        time_s = now.strftime("%H:%M:%S")
        vnum = vehicle_number or "Unknown"

        with self._lock:
            try:
                wb = openpyxl.load_workbook(str(self.filepath))

                # 1 ── All Detections ─────────────────────────────────────────
                ws_all = wb[SH_ALL]
                ws_all.append([date_s, time_s, vnum,
                               "YES" if logo1 else "NO",
                               "YES" if logo2 else "NO",
                               "YES" if logo3 else "NO",
                               count])
                r = ws_all.max_row
                f = _FILLS.get(count, _FILLS[0])
                for c in range(1, len(_ALL_HDR) + 1):
                    ws_all.cell(r, c).fill = f

                # 2 ── Should Install (stickers < 3) ─────────────────────────
                if count < 3:
                    missing = [n for flag, n in [
                        (logo1, "Triangle Logo"),
                        (logo2, "A.M.MOTORS text"),
                        (logo3, "Website sticker"),
                    ] if not flag]
                    action = "Install: " + ", ".join(missing)
                    ws_inst = wb[SH_INSTALL]
                    ws_inst.append([date_s, time_s, vnum,
                                   "YES" if logo1 else "NO",
                                   "YES" if logo2 else "NO",
                                   "YES" if logo3 else "NO",
                                   3 - count, action])
                    ri = ws_inst.max_row
                    for c in range(1, len(_INSTALL_HDR) + 1):
                        ws_inst.cell(ri, c).fill = f

                wb.save(str(self.filepath))
                return True
            except Exception as e:
                print(f"[Logger] log error: {e}")
                return False

    def log_unknown(
        self,
        vehicle_number: str,
        frame: Optional[np.ndarray],
        notes: str = "Unknown / non-AM-Motors sticker detected",
    ) -> bool:
        now = datetime.now()
        with self._lock:
            try:
                wb = openpyxl.load_workbook(str(self.filepath))
                ws = wb[SH_UNKNOWN]
                ws.append([now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S"),
                           vehicle_number or "Unknown", notes, ""])
                rn = ws.max_row

                xl_img = _frame_to_xl_image(frame) if frame is not None else None
                if xl_img is not None:
                    ws.add_image(xl_img, f"E{rn}")
                    ws.row_dimensions[rn].height = 98

                wb.save(str(self.filepath))
                return True
            except Exception as e:
                print(f"[Logger] unknown sticker log error: {e}")
                return False

    # ── Analytics ─────────────────────────────────────────────────────────────

    def _all_rows(self) -> list[dict]:
        try:
            wb = openpyxl.load_workbook(str(self.filepath), read_only=True, data_only=True)
            ws = wb[SH_ALL]
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
        except Exception:
            return []
        if len(rows) <= 1:
            return []
        hdrs = rows[0]
        return [dict(zip(hdrs, r)) for r in rows[1:] if any(r)]

    def get_period_stats(self, period: str = "today", reach_factor: int = 50) -> dict:
        """
        period: "today" | "yesterday" | "week" | "month" | "year" | "all"
        Returns dict: total, full, partial, none, reach
        """
        rows = self._all_rows()
        today = datetime.now().date()

        def _in_period(d_str):
            try:
                d = datetime.strptime(str(d_str), "%Y-%m-%d").date()
            except Exception:
                return False
            if period == "today":
                return d == today
            if period == "yesterday":
                return d == today - timedelta(days=1)
            if period == "week":
                return (today - d).days < 7
            if period == "month":
                return d.year == today.year and d.month == today.month
            if period == "year":
                return d.year == today.year
            return True  # "all"

        filtered = [r for r in rows if _in_period(r.get("Date"))]

        def _count(r):
            try:
                return int(r.get("Stickers_Found", 0) or 0)
            except Exception:
                return 0

        total   = len(filtered)
        full    = sum(1 for r in filtered if _count(r) == 3)
        partial = sum(1 for r in filtered if 1 <= _count(r) < 3)
        none    = sum(1 for r in filtered if _count(r) == 0)
        reach   = (full + partial) * max(1, reach_factor)

        return dict(total=total, full=full, partial=partial, none=none, reach=reach)

    def get_daily_series(self, days: int = 30) -> dict:
        """Return {date_str: vehicle_count} for the last N calendar days."""
        rows = self._all_rows()
        today = datetime.now().date()
        series = {
            (today - timedelta(days=i)).strftime("%Y-%m-%d"): 0
            for i in range(days - 1, -1, -1)
        }
        for r in rows:
            d = str(r.get("Date", ""))
            if d in series:
                series[d] += 1
        return series

    def get_sticker_breakdown(self, period: str = "today") -> dict:
        """Return count for each combination: {0:n, 1:n, 2:n, 3:n}."""
        rows = self._all_rows()
        today = datetime.now().date()

        def _in(d_str):
            try:
                d = datetime.strptime(str(d_str), "%Y-%m-%d").date()
            except Exception:
                return False
            if period == "today":      return d == today
            if period == "yesterday":  return d == today - timedelta(days=1)
            if period == "week":       return (today - d).days < 7
            if period == "month":      return d.year == today.year and d.month == today.month
            if period == "year":       return d.year == today.year
            return True

        out = {0: 0, 1: 0, 2: 0, 3: 0}
        for r in rows:
            if _in(r.get("Date")):
                try:
                    k = int(r.get("Stickers_Found", 0) or 0)
                    out[k] = out.get(k, 0) + 1
                except Exception:
                    pass
        return out

    # ── Simple reads ──────────────────────────────────────────────────────────

    def get_recent(self, n: int = 15, sheet: str = SH_ALL) -> list[dict]:
        try:
            wb = openpyxl.load_workbook(str(self.filepath), read_only=True, data_only=True)
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
        except Exception:
            return []
        if len(rows) <= 1:
            return []
        hdrs = rows[0]
        return [dict(zip(hdrs, r)) for r in rows[max(1, len(rows) - n):]]

    @property
    def total_logged(self) -> int:
        try:
            wb = openpyxl.load_workbook(str(self.filepath), read_only=True)
            n = wb[SH_ALL].max_row - 1
            wb.close()
            return max(0, n)
        except Exception:
            return 0
